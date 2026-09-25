import openai
import base64
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# --- CONFIGURACIÓN ---
client = openai.OpenAI(api_key="TU_LLAVE_AQUI")

def encode_image(image_path):
    with open(image_path, "rb") as image_file:
        return base64.b64encode(image_file.read()).decode('utf-8')

def extraer_temas_de_imagen(ruta_imagen):
    base64_image = encode_image(ruta_imagen)
    prompt_vision = (
        "Analiza la imagen e identifica los temas. "
        "Si un punto tiene sub-elementos (ej: Gastronomía, Lengua), devuélvelos como temas individuales. "
        "Devuelve solo los nombres de los temas separados por comas."
    )
    response = client.chat.completions.create(
        model="gpt-4o",
        messages=[{"role": "user", "content": [
            {"type": "text", "text": prompt_vision},
            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}}
        ]}]
    )
    return [tema.strip() for tema in response.choices[0].message.content.split(',') if tema.strip()]

def generar_informe_final(ruta_imagen):
    temas_extraidos = extraer_temas_de_imagen(ruta_imagen)
    # Filtro de seguridad
    temas_finales = [t for t in temas_extraidos if not any(p in t.lower() for p in ["soberanía", "4to año", "20/04"])]

    if not temas_finales: return
    tema_contexto = temas_finales[0]
    doc = Document()
    
    for i, tema in enumerate(temas_finales):
        numero = i + 1
        print(f"[{numero}/{len(temas_finales)}] Redactando: {tema}...")
        
        # Títulos limpios
        t_clean = tema.replace("-", "").replace("•", "").strip().strip('¿?').capitalize()
        es_p = t_clean.lower().startswith(("que", "como", "cual", "por que", "por qué", "quien", "donde"))
        titulo_final = f"{numero}. ¿{t_clean}?" if es_p else f"{numero}. {t_clean}"

        # --- INSTRUCCIONES DE REDACCIÓN BLINDADAS ---
        instrucciones_redaccion = (
            f"Actúa como experto en '{tema_contexto}'. Redacta sobre '{tema}'. "
            "ESTRUCTURA OBLIGATORIA:\n"
            "1. Empieza con un párrafo académico formal de al menos 70 palabras.\n"
            "2. SOLO AL FINAL, si se pide lista, añade máximo 4 viñetas cortas.\n"
            "REGLA CRÍTICA: No hagas que todo el texto sea una lista. No uses etiquetas como 'componente:'. "
            "Extensión total: 115-125 palabras. Todo en minúsculas (excepto inicio de oración). Sin ';' ni '**'."
        )

        usa_lista = " Al final incluye una lista de 3 viñetas '• Concepto: breve descripción'." if i % 2 == 0 else " NO uses viñetas, solo párrafos."
        
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "system", "content": instrucciones_redaccion},
                {"role": "user", "content": f"Desarrolle '{tema}' en el contexto de '{tema_contexto}'.{usa_lista}"}
            ],
            temperature=0.4
        )
        
        # Limpieza de símbolos basura
        limpio = response.choices[0].message.content.replace("••", "").replace("componente:", "").replace("*", "")
        limpio = limpio.replace("\n- ", "\n• ").replace(". •", ".\n•")

        # --- WORD ---
        h = doc.add_paragraph()
        run_h = h.add_run(titulo_final)
        run_h.bold = True
        run_h.font.name = 'Canva Sans'
        run_h.font.size = Pt(15)

        lineas = limpio.split('\n')
        for linea in lineas:
            linea = linea.strip()
            if not linea: continue
            
            # Si la línea es muy larga (párrafo), la dividimos si es necesario
            if not linea.startswith("•") and len(linea.split()) > 65:
                frases = linea.split('. ')
                m = len(frases) // 2
                bloques = [". ".join(frases[:m]) + ".", ". ".join(frases[m:])] if m > 0 else [linea]
            else:
                bloques = [linea]

            for bloque in bloques:
                p = doc.add_paragraph()
                if bloque.strip().startswith("•"):
                    p.add_run("• ").bold = True
                    cont = bloque.strip().lstrip("• ").strip()
                    if ":" in cont:
                        sub, desc = cont.split(":", 1)
                        run_s = p.add_run(f"{sub.strip().capitalize()}:")
                        run_s.bold = True
                        p.add_run(desc)
                    else:
                        p.add_run(cont)
                    p.paragraph_format.left_indent = Pt(24)
                else:
                    p.add_run(bloque.strip())
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                
                for run in p.runs:
                    run.font.name = 'Canva Sans'
                    run.font.size = Pt(13)

        doc.add_paragraph()

    doc.save("Informe_Final_Pro_Perfecto.docx")
    print(f"\n¡Listo! Procesados {len(temas_finales)} temas con equilibrio párrafo/lista.")

if __name__ == "__main__":
    ruta = input("Ingresa la ruta de la imagen: ").strip('"')
    generar_informe_final(ruta)