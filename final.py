import openai
import base64
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# --- CONFIGURACIÓN ---
client = openai.OpenAI(api_key="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")

def encode_image(image_path):
    with open(image_path, "rb") as image_file:
        return base64.b64encode(image_file.read()).decode('utf-8')

def extraer_temas_de_imagen(ruta_imagen):
    print(f"--- Analizando imagen detalladamente: {ruta_imagen} ---")
    base64_image = encode_image(ruta_imagen)
    
    # PROMPT REFORZADO PARA DESGLOSAR SUBPUNTOS
    prompt_vision = (
        "Analiza esta imagen escolar. "
        "1. Identifica el TEMA GLOBAL (ej. Manifestaciones culturales indígenas). "
        "2. Identifica todos los puntos con viñetas o asteriscos. "
        "REGLA CRÍTICA: Si un punto menciona 'Principales manifestaciones' y tiene una lista debajo "
        "(como Gastronomía, Lengua, Medicina, Artesanía, etc.), DEBES devolver cada uno de esos "
        "elementos como un tema individual separado por comas. "
        "No resumas. Necesito que cada subtema sea un punto independiente en la lista final. "
        "Devuelve solo los nombres de los temas separados por comas, sin etiquetas."
    )

    response = client.chat.completions.create(
        model="gpt-4o",
        messages=[{"role": "user", "content": [
            {"type": "text", "text": prompt_vision},
            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}}
        ]}]
    )
    
    contenido = response.choices[0].message.content
    for etiqueta in ["Tema principal:", "Tema global:", "Puntos:", "Subpuntos:"]:
        contenido = contenido.replace(etiqueta, "")
        
    return [tema.strip() for tema in contenido.split(',') if tema.strip()]

def generar_informe_final(ruta_imagen):
    temas_extraidos = extraer_temas_de_imagen(ruta_imagen)
    
    # Filtro de ruido
    palabras_prohibidas = ["soberanía", "4to año", "lunes", "investigación", "20/04"]
    temas_finales = [t for t in temas_extraidos if not any(p in t.lower() for p in palabras_prohibidas)]

    if not temas_finales:
        print("No se detectaron temas.")
        return

    # Usamos el primer tema real como contexto (ej: Manifestaciones culturales)
    tema_contexto = temas_finales[0]
    doc = Document()
    
    print(f"Temas detectados para redactar: {len(temas_finales)}")

    for i, tema in enumerate(temas_finales):
        numero = i + 1
        print(f"[{numero}/{len(temas_finales)}] Redactando: {tema}...")
        
        # --- LIMPIEZA Y FORMATEO DE TÍTULOS ---
        t_clean = tema.replace("-", "-").replace("•", "").strip()
        t_clean = t_clean.strip('¿?').capitalize()
        t_low = t_clean.lower()
        
        es_pregunta = t_low.startswith(("que", "como", "cual", "por que", "por qué", "quien", "donde"))
        titulo_final = f"{numero}. ¿{t_clean}?" if es_pregunta else f"{numero}. {t_clean}"

        # Instrucciones de redacción (Contexto invisible aplicado)
        instrucciones_redaccion = (
            f"Eres un experto en '{tema_contexto}'. Desarrolla el subtema con rigor académico. "
            "REGLA CRÍTICA: El texto completo debe tener entre 115 y 125 palabras. "
            "No incluyas títulos. Empieza directo. Prohibido usar punto y coma (;). "
            "MINÚSCULAS: Todo en minúsculas, EXCEPTO la primera letra de cada oración y la primera letra en titulos o encabezados. "
            "LISTAS: Usa •. Formato 'Componente: descripción breve' (máximo 20 palabras por punto)."
        )

        # Alternamos formato para que no sea repetitivo
        formato_lista = " DEBES incluir una lista formal 'Componente: descripción'." if i % 2 == 0 else " No uses listas."
        
        response = client.chat.completions.create(
            model="gpt-4o",
            messages=[
                {"role": "system", "content": instrucciones_redaccion},
                {"role": "user", "content": f"Desarrolle el punto '{tema}' enfocado en '{tema_contexto}'.{formato_lista}"}
            ],
            temperature=0.4
        )
        
        texto_generado = response.choices[0].message.content.replace("*", "").replace(";", ".")

        # --- FORMATO WORD ---
        h = doc.add_paragraph()
        run_h = h.add_run(titulo_final)
        run_h.bold = True
        run_h.font.name = 'Canva Sans'
        run_h.font.size = Pt(15)

        lineas = texto_generado.split('\n')
        for linea in lineas:
            linea = linea.strip()
            if not linea: continue
            
            # SEPARACIÓN DE PÁRRAFOS LARGOS
            if not linea.startswith("•") and len(linea.split()) > 65:
                puntos = linea.split('. ')
                mitad = len(puntos) // 2
                bloques = [". ".join(puntos[:mitad]) + ".", ". ".join(puntos[mitad:])] if mitad > 0 else [linea]
            else:
                bloques = [linea]

            for bloque in bloques:
                p = doc.add_paragraph()
                if bloque.strip().startswith("•"):
                    p.add_run("• ").bold = True
                    cont = bloque.strip().lstrip("• ").strip()
                    if ":" in cont:
                        sub, desc = cont.split(":", 1)
                        run_sub = p.add_run(f"{sub.strip()}:")
                        run_sub.bold = True
                        p.add_run(desc)
                    else:
                        p.add_run(cont)
                    p.paragraph_format.left_indent = Pt(24)
                else:
                    p.add_run(bloque.strip())
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                
                for run in p.runs:
                    run.font.name = 'Canva Sans'
                    run.font.size = Pt(13)

        doc.add_paragraph()

    doc.save("Informe_Completo_9Temas.docx")
    print(f"\n¡Listo! Se procesaron {len(temas_finales)} temas individuales.")

if __name__ == "__main__":
    ruta = input("Ingresa la ruta de la imagen: ").strip('"')
    generar_informe_final(ruta)

    #Necesito hacer que sea consistente las listas y que no uqede una lista en forma de parrafo 19/04/26